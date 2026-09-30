import csv
import io
import os
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import (Flask,g,jsonify,redirect,render_template,request,send_file,session,url_for, )

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import logica


app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "cambia-esta-clave"
)

BD = os.environ.get(
    "BD_ASISTENCIA",
    "asistencia.db"
)

CLAVE_DOCENTE = os.environ.get(
    "CLAVE_DOCENTE",
    "docente123"
)

SEGUNDOS_SIN_SENAL = int(
    os.environ.get(
        "SEGUNDOS_SIN_SENAL",
        "15"
    )
)

ZONA_HORARIA = ZoneInfo("America/Lima")


# ============================================================
# HORA
# ============================================================

def hora_peru():
    """Hora actual de Perú."""
    return datetime.now(
        ZONA_HORARIA
    ).replace(tzinfo=None)


# ============================================================
# BASE DE DATOS
# ============================================================

def bd():
    """Obtiene una conexión SQLite para la petición actual."""

    if "bd" not in g:
        g.bd = sqlite3.connect(
            BD,
            timeout=30
        )
        g.bd.row_factory = sqlite3.Row

    return g.bd


@app.teardown_appcontext
def cerrar_bd(_):
    conexion = g.pop("bd", None)

    if conexion:
        conexion.close()


def crear_tablas():
    """Crea las tablas y carga los estudiantes."""

    con = sqlite3.connect(
        BD,
        timeout=30
    )

    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS estudiantes (
            codigo TEXT PRIMARY KEY,
            nombre TEXT
        );

        CREATE TABLE IF NOT EXISTS clases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            curso TEXT,
            hora_inicio TEXT,
            hora_fin TEXT,
            tolerancia_min INTEGER,
            activa INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS registros (
            clase_id INTEGER,
            codigo TEXT,
            hora_ingreso TEXT,
            estado_ingreso TEXT,
            desconexiones INTEGER,
            desconectado_desde TEXT,
            ultima_senal TEXT,
            retiro INTEGER,
            PRIMARY KEY (clase_id, codigo)
        );

        CREATE TABLE IF NOT EXISTS eventos (
            clase_id INTEGER,
            codigo TEXT,
            tipo TEXT,
            hora TEXT
        );
        """
    )

    # --------------------------------------------------------
    # Cargar estudiantes.csv
    # --------------------------------------------------------

    if os.path.exists("estudiantes.csv"):

        with open(
            "estudiantes.csv",
            encoding="utf-8"
        ) as f:

            filas = []

            for r in csv.DictReader(f):

                codigo = (
                    r.get("codigo", "")
                    .strip()
                    .upper()
                )

                nombre = (
                    r.get("nombre", "")
                    .strip()
                )

                if codigo:
                    filas.append(
                        (
                            codigo,
                            nombre
                        )
                    )

            con.executemany(
                """
                INSERT OR REPLACE INTO estudiantes
                (codigo, nombre)
                VALUES (?, ?)
                """,
                filas
            )

    con.commit()
    con.close()


# ============================================================
# CLASE
# ============================================================

def clase_activa():
    return bd().execute(
        """
        SELECT *
        FROM clases
        WHERE activa = 1
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()


# ============================================================
# REGISTROS
# ============================================================

def leer_registro(clase_id, codigo):
    """Busca el registro de un alumno."""

    fila = bd().execute(
        """
        SELECT *
        FROM registros
        WHERE clase_id = ?
        AND codigo = ?
        """,
        (
            clase_id,
            codigo
        )
    ).fetchone()

    return dict(fila) if fila else None


def guardar_registro(
    clase_id,
    reg,
    eventos
):
    """Guarda el registro y sus eventos."""

    bd().execute(
        """
        INSERT OR REPLACE INTO registros
        (
            clase_id,
            codigo,
            hora_ingreso,
            estado_ingreso,
            desconexiones,
            desconectado_desde,
            ultima_senal,
            retiro
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            clase_id,
            reg["codigo"],
            reg["hora_ingreso"],
            reg["estado_ingreso"],
            reg["desconexiones"],
            reg["desconectado_desde"],
            reg["ultima_senal"],
            reg["retiro"]
        )
    )

    for tipo, hora in eventos:

        bd().execute(
            """
            INSERT INTO eventos
            (
                clase_id,
                codigo,
                tipo,
                hora
            )
            VALUES (?, ?, ?, ?)
            """,
            (
                clase_id,
                reg["codigo"],
                tipo,
                hora.isoformat(
                    timespec="seconds"
                )
            )
        )

    bd().commit()


# ============================================================
# MONITOREO
# ============================================================

def revisar_clase(clase_id):
    """
    Revisa los alumnos que ya ingresaron
    y detecta pérdidas de conexión.
    """

    ahora = hora_peru()

    filas = bd().execute(
        """
        SELECT *
        FROM registros
        WHERE clase_id = ?
        AND retiro = 0
        """,
        (clase_id,)
    ).fetchall()

    for fila in filas:

        reg = dict(fila)
        eventos = []

        logica.revisar_tiempos(
            reg,
            ahora,
            SEGUNDOS_SIN_SENAL,
            eventos
        )

        guardar_registro(
            clase_id,
            reg,
            eventos
        )


# ============================================================
# INGRESO DEL ALUMNO
# ============================================================

@app.route("/", methods=["GET", "POST"])
def ingreso():

    clase = clase_activa()
    mensaje = None

    if request.method == "POST":

        codigo = (
            request.form
            .get("codigo", "")
            .strip()
            .upper()
        )

        estudiante = bd().execute(
            """
            SELECT *
            FROM estudiantes
            WHERE codigo = ?
            """,
            (codigo,)
        ).fetchone()

        if not clase:

            mensaje = (
                "La clase todavía no ha iniciado. "
                "Espera a que el docente la inicie."
            )

        elif not estudiante:

            mensaje = (
                "Código inválido. "
                "Revisa tu código institucional."
            )

        else:

            ahora = hora_peru()

            reg = leer_registro(
                clase["id"],
                codigo
            )

            eventos = []

            # ------------------------------------------------
            # PRIMER INGRESO
            # ------------------------------------------------

            if reg is None:

                reg = logica.nuevo_registro(
                    codigo,
                    ahora,
                    logica.a_fecha(
                        clase["hora_inicio"]
                    ),
                    clase["tolerancia_min"]
                )

            # ------------------------------------------------
            # VOLVIÓ A ENTRAR
            # ------------------------------------------------

            else:

                logica.revisar_tiempos(
                    reg,
                    ahora,
                    SEGUNDOS_SIN_SENAL,
                    eventos
                )

                if (
                    reg["desconectado_desde"]
                    and not reg["retiro"]
                ):

                    logica.registrar_reingreso(
                        reg,
                        ahora,
                        eventos
                    )

                if not reg["retiro"]:

                    reg["ultima_senal"] = (
                        ahora.isoformat(
                            timespec="seconds"
                        )
                    )

            guardar_registro(
                clase["id"],
                reg,
                eventos
            )

            session["codigo"] = codigo

            return redirect(
                url_for("aula")
            )

    return render_template(
        "ingreso.html",
        clase=clase,
        mensaje=mensaje
    )


# ============================================================
# AULA DEL ALUMNO
# ============================================================

@app.route("/aula")
def aula():

    clase = clase_activa()
    codigo = session.get("codigo")

    if not clase or not codigo:
        return redirect(
            url_for("ingreso")
        )

    reg = leer_registro(
        clase["id"],
        codigo
    )

    if reg is None:
        return redirect(
            url_for("ingreso")
        )

    estudiante = bd().execute(
        """
        SELECT *
        FROM estudiantes
        WHERE codigo = ?
        """,
        (codigo,)
    ).fetchone()

    return render_template(
        "aula.html",
        clase=clase,
        reg=reg,
        estudiante=estudiante,
        estado=logica.estado_final(reg)
    )


# ============================================================
# SEÑAL DEL NAVEGADOR
# ============================================================

@app.route("/api/senal", methods=["POST"])
def senal():

    clase = clase_activa()
    codigo = session.get("codigo")

    if not clase or not codigo:

        return jsonify(
            clase_activa=False
        )

    reg = leer_registro(
        clase["id"],
        codigo
    )

    if reg is None:

        return jsonify(
            clase_activa=False
        )

    datos = request.get_json(
        silent=True
    ) or {}

    ahora = hora_peru()
    eventos = []

    visible = datos.get(
        "visible",
        True
    )

    # --------------------------------------------------------
    # Comprobar si había una desconexión
    # --------------------------------------------------------

    logica.revisar_tiempos(
        reg,
        ahora,
        SEGUNDOS_SIN_SENAL,
        eventos
    )

    # --------------------------------------------------------
    # PÁGINA VISIBLE
    # --------------------------------------------------------

    if visible:

        if (
            reg["desconectado_desde"]
            and not reg["retiro"]
        ):

            logica.registrar_reingreso(
                reg,
                ahora,
                eventos
            )

        if not reg["retiro"]:

            reg["ultima_senal"] = (
                ahora.isoformat(
                    timespec="seconds"
                )
            )

    # --------------------------------------------------------
    # PÁGINA NO VISIBLE
    # --------------------------------------------------------

    else:

        logica.registrar_salida(
            reg,
            ahora,
            eventos
        )

        if not reg["retiro"]:

            reg["ultima_senal"] = (
                ahora.isoformat(
                    timespec="seconds"
                )
            )

    guardar_registro(
        clase["id"],
        reg,
        eventos
    )

    return jsonify(
        clase_activa=True,
        estado=logica.estado_final(reg),
        desconexiones=reg["desconexiones"]
    )


# ============================================================
# PANEL DOCENTE
# ============================================================

@app.route("/docente", methods=["GET", "POST"])
def docente():

    if (
        request.method == "POST"
        and "clave" in request.form
    ):

        if request.form["clave"] == CLAVE_DOCENTE:

            session["docente"] = True

        else:

            return render_template(
                "docente_login.html",
                mensaje="Clave incorrecta."
            )

    if not session.get("docente"):

        return render_template(
            "docente_login.html",
            mensaje=None
        )

    clase = clase_activa()
    filas = []

    if clase:

        revisar_clase(
            clase["id"]
        )

        filas = filas_reporte(
            clase["id"]
        )

    ultima = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE activa = 0
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()

    return render_template(
        "docente.html",
        clase=clase,
        filas=filas,
        ultima=ultima
    )


# ============================================================
# INICIAR CLASE
# ============================================================

@app.post("/docente/iniciar")
def iniciar_clase():

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    if not clase_activa():

        curso = (
            request.form.get("curso")
            or "Clase virtual"
        )

        tolerancia = int(
            request.form.get(
                "tolerancia"
            ) or 10
        )

        bd().execute(
            """
            INSERT INTO clases
            (
                curso,
                hora_inicio,
                tolerancia_min
            )
            VALUES (?, ?, ?)
            """,
            (
                curso,
                hora_peru().isoformat(
                    timespec="seconds"
                ),
                tolerancia
            )
        )

        bd().commit()

    return redirect(
        url_for("docente")
    )


# ============================================================
# FINALIZAR CLASE
# ============================================================

@app.post("/docente/finalizar")
def finalizar_clase():

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = clase_activa()

    if clase:

        revisar_clase(
            clase["id"]
        )

        bd().execute(
            """
            UPDATE clases
            SET activa = 0,
                hora_fin = ?
            WHERE id = ?
            """,
            (
                hora_peru().isoformat(
                    timespec="seconds"
                ),
                clase["id"]
            )
        )

        bd().commit()

        return redirect(
            url_for(
                "reporte",
                clase_id=clase["id"]
            )
        )

    return redirect(
        url_for("docente")
    )


# ============================================================
# GENERAR FILAS DEL REPORTE
# ============================================================

def filas_reporte(clase_id):

    """
    IMPORTANTE:

    Se recorren TODOS los estudiantes.

    Si el alumno nunca ingresó:
        estado_final = FALTA

    Si ingresó:
        se utiliza su registro real.
    """

    filas = []

    estudiantes = bd().execute(
        """
        SELECT *
        FROM estudiantes
        ORDER BY nombre
        """
    ).fetchall()

    for est in estudiantes:

        codigo = est["codigo"]

        reg = leer_registro(
            clase_id,
            codigo
        )

        # ====================================================
        # NUNCA INGRESÓ
        # ====================================================

        if reg is None:

            filas.append(
                {
                    "codigo": codigo,
                    "nombre": est["nombre"],
                    "hora_ingreso": "",
                    "estado_ingreso": "",
                    "desconexiones": 0,
                    "salidas": "",
                    "reingresos": "",
                    "conectado": False,
                    "estado_final": logica.FALTA
                }
            )

            continue

        # ====================================================
        # SÍ INGRESÓ
        # ====================================================

        eventos = bd().execute(
            """
            SELECT tipo, hora
            FROM eventos
            WHERE clase_id = ?
            AND codigo = ?
            ORDER BY hora
            """,
            (
                clase_id,
                codigo
            )
        ).fetchall()

        salidas = ", ".join(
            e["hora"][11:]
            for e in eventos
            if e["tipo"] == "salida"
        )

        reingresos = ", ".join(
            e["hora"][11:]
            for e in eventos
            if e["tipo"] == "reingreso"
        )

        conectado = (
            not reg["desconectado_desde"]
            and not reg["retiro"]
        )

        filas.append(
            {
                "codigo": codigo,
                "nombre": est["nombre"],
                "hora_ingreso": (
                    reg["hora_ingreso"][11:]
                    if reg["hora_ingreso"]
                    else ""
                ),
                "estado_ingreso": (
                    reg["estado_ingreso"]
                    or ""
                ),
                "desconexiones": (
                    reg["desconexiones"]
                    or 0
                ),
                "salidas": salidas,
                "reingresos": reingresos,
                "conectado": conectado,
                "estado_final": logica.estado_final(reg)
            }
        )

    return filas


# ============================================================
# REPORTE HTML
# ============================================================

@app.route("/reporte/<int:clase_id>")
def reporte(clase_id):

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE id = ?
        """,
        (clase_id,)
    ).fetchone()

    if not clase:

        return redirect(
            url_for("docente")
        )

    revisar_clase(
        clase_id
    )

    filas = filas_reporte(
        clase_id
    )

    estados = (
        logica.PRESENTE,
        logica.TARDANZA,
        logica.FALTA,
        logica.RETIRO_FALTA
    )

    resumen = {
        estado: sum(
            1
            for f in filas
            if f["estado_final"] == estado
        )
        for estado in estados
    }

    return render_template(
        "reporte.html",
        clase=clase,
        filas=filas,
        resumen=resumen
    )


# ============================================================
# EXPORTAR A EXCEL
# ============================================================

@app.route("/reporte/<int:clase_id>/excel")
def reporte_excel(clase_id):

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE id = ?
        """,
        (clase_id,)
    ).fetchone()

    if not clase:

        return redirect(
            url_for("docente")
        )

    revisar_clase(
        clase_id
    )

    filas = filas_reporte(
        clase_id
    )

    colores = {
        logica.PRESENTE: "C8E6D4",
        logica.TARDANZA: "F6E2B3",
        logica.FALTA: "F2C4C0",
        logica.RETIRO_FALTA: "DCCDEE"
    }

    libro = Workbook()

    hoja = libro.active
    hoja.title = "Asistencia"

    # --------------------------------------------------------
    # ENCABEZADO
    # --------------------------------------------------------

    hoja.append(
        [
            f"Reporte de asistencia - "
            f"{clase['curso']}"
        ]
    )

    hoja.append(
        [
            f"Fecha: {clase['hora_inicio'][:10]}  "
            f"Inicio: {clase['hora_inicio'][11:]}  "
            f"Fin: {(clase['hora_fin'] or '')[11:]}"
        ]
    )

    hoja.append([])

    hoja.append(
        [
            "Código",
            "Estudiante",
            "Hora de ingreso",
            "Estado de ingreso",
            "Desconexiones",
            "Horas de salida",
            "Horas de reingreso",
            "Estado final"
        ]
    )

    # --------------------------------------------------------
    # FORMATO DEL ENCABEZADO
    # --------------------------------------------------------

    for celda in hoja[4]:

        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celda.fill = PatternFill(
            "solid",
            fgColor="1F3A5F"
        )

    # --------------------------------------------------------
    # DATOS
    # --------------------------------------------------------

    for f in filas:

        estado = f["estado_final"]

        hoja.append(
            [
                f["codigo"],
                f["nombre"],
                f["hora_ingreso"],
                f["estado_ingreso"],
                f["desconexiones"],
                f["salidas"],
                f["reingresos"],
                estado
            ]
        )

        # Color según estado
        hoja.cell(
            hoja.max_row,
            8
        ).fill = PatternFill(
            "solid",
            fgColor=colores.get(
                estado,
                "FFFFFF"
            )
        )

    # --------------------------------------------------------
    # ANCHOS
    # --------------------------------------------------------

    hoja["A1"].font = Font(
        bold=True,
        size=14
    )

    anchos = (
        12,
        32,
        15,
        17,
        14,
        24,
        24,
        16
    )

    for col, ancho in zip(
        "ABCDEFGH",
        anchos
    ):

        hoja.column_dimensions[
            col
        ].width = ancho

    # --------------------------------------------------------
    # CREAR ARCHIVO
    # --------------------------------------------------------

    archivo = io.BytesIO()

    libro.save(
        archivo
    )

    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name=(
            f"asistencia_"
            f"{clase['hora_inicio'][:10]}_"
            f"clase{clase_id}.xlsx"
        )
    )


# ============================================================
# SALIR
# ============================================================

@app.route("/salir")
def salir():

    session.clear()

    return redirect(
        url_for("ingreso")
    )


# ============================================================
# CREAR BASE DE DATOS
# ============================================================

crear_tablas()


# ============================================================
# EJECUTAR
# ============================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )
import csv
import io
import os
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import (
    Flask,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import logica


app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "cambia-esta-clave"
)

BD = os.environ.get(
    "BD_ASISTENCIA",
    "asistencia.db"
)

CLAVE_DOCENTE = os.environ.get(
    "CLAVE_DOCENTE",
    "docente123"
)

SEGUNDOS_SIN_SENAL = int(
    os.environ.get(
        "SEGUNDOS_SIN_SENAL",
        "15"
    )
)

ZONA_HORARIA = ZoneInfo("America/Lima")


# ============================================================
# HORA
# ============================================================

def hora_peru():
    """Hora actual de Perú."""
    return datetime.now(
        ZONA_HORARIA
    ).replace(tzinfo=None)


# ============================================================
# BASE DE DATOS
# ============================================================

def bd():
    """Obtiene una conexión SQLite para la petición actual."""

    if "bd" not in g:
        g.bd = sqlite3.connect(
            BD,
            timeout=30
        )
        g.bd.row_factory = sqlite3.Row

    return g.bd


@app.teardown_appcontext
def cerrar_bd(_):
    conexion = g.pop("bd", None)

    if conexion:
        conexion.close()


def crear_tablas():
    """Crea las tablas y carga los estudiantes."""

    con = sqlite3.connect(
        BD,
        timeout=30
    )

    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS estudiantes (
            codigo TEXT PRIMARY KEY,
            nombre TEXT
        );

        CREATE TABLE IF NOT EXISTS clases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            curso TEXT,
            hora_inicio TEXT,
            hora_fin TEXT,
            tolerancia_min INTEGER,
            activa INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS registros (
            clase_id INTEGER,
            codigo TEXT,
            hora_ingreso TEXT,
            estado_ingreso TEXT,
            desconexiones INTEGER,
            desconectado_desde TEXT,
            ultima_senal TEXT,
            retiro INTEGER,
            PRIMARY KEY (clase_id, codigo)
        );

        CREATE TABLE IF NOT EXISTS eventos (
            clase_id INTEGER,
            codigo TEXT,
            tipo TEXT,
            hora TEXT
        );
        """
    )

    # --------------------------------------------------------
    # Cargar estudiantes.csv
    # --------------------------------------------------------

    if os.path.exists("estudiantes.csv"):

        with open(
            "estudiantes.csv",
            encoding="utf-8"
        ) as f:

            filas = []

            for r in csv.DictReader(f):

                codigo = (
                    r.get("codigo", "")
                    .strip()
                    .upper()
                )

                nombre = (
                    r.get("nombre", "")
                    .strip()
                )

                if codigo:
                    filas.append(
                        (
                            codigo,
                            nombre
                        )
                    )

            con.executemany(
                """
                INSERT OR REPLACE INTO estudiantes
                (codigo, nombre)
                VALUES (?, ?)
                """,
                filas
            )

    con.commit()
    con.close()


# ============================================================
# CLASE
# ============================================================

def clase_activa():
    return bd().execute(
        """
        SELECT *
        FROM clases
        WHERE activa = 1
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()


# ============================================================
# REGISTROS
# ============================================================

def leer_registro(clase_id, codigo):
    """Busca el registro de un alumno."""

    fila = bd().execute(
        """
        SELECT *
        FROM registros
        WHERE clase_id = ?
        AND codigo = ?
        """,
        (
            clase_id,
            codigo
        )
    ).fetchone()

    return dict(fila) if fila else None


def guardar_registro(
    clase_id,
    reg,
    eventos
):
    """Guarda el registro y sus eventos."""

    bd().execute(
        """
        INSERT OR REPLACE INTO registros
        (
            clase_id,
            codigo,
            hora_ingreso,
            estado_ingreso,
            desconexiones,
            desconectado_desde,
            ultima_senal,
            retiro
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            clase_id,
            reg["codigo"],
            reg["hora_ingreso"],
            reg["estado_ingreso"],
            reg["desconexiones"],
            reg["desconectado_desde"],
            reg["ultima_senal"],
            reg["retiro"]
        )
    )

    for tipo, hora in eventos:

        bd().execute(
            """
            INSERT INTO eventos
            (
                clase_id,
                codigo,
                tipo,
                hora
            )
            VALUES (?, ?, ?, ?)
            """,
            (
                clase_id,
                reg["codigo"],
                tipo,
                hora.isoformat(
                    timespec="seconds"
                )
            )
        )

    bd().commit()


# ============================================================
# MONITOREO
# ============================================================

def revisar_clase(clase_id):
    """
    Revisa los alumnos que ya ingresaron
    y detecta pérdidas de conexión.
    """

    ahora = hora_peru()

    filas = bd().execute(
        """
        SELECT *
        FROM registros
        WHERE clase_id = ?
        AND retiro = 0
        """,
        (clase_id,)
    ).fetchall()

    for fila in filas:

        reg = dict(fila)
        eventos = []

        logica.revisar_tiempos(
            reg,
            ahora,
            SEGUNDOS_SIN_SENAL,
            eventos
        )

        guardar_registro(
            clase_id,
            reg,
            eventos
        )


# ============================================================
# INGRESO DEL ALUMNO
# ============================================================

@app.route("/", methods=["GET", "POST"])
def ingreso():

    clase = clase_activa()
    mensaje = None

    if request.method == "POST":

        codigo = (
            request.form
            .get("codigo", "")
            .strip()
            .upper()
        )

        estudiante = bd().execute(
            """
            SELECT *
            FROM estudiantes
            WHERE codigo = ?
            """,
            (codigo,)
        ).fetchone()

        if not clase:

            mensaje = (
                "La clase todavía no ha iniciado. "
                "Espera a que el docente la inicie."
            )

        elif not estudiante:

            mensaje = (
                "Código inválido. "
                "Revisa tu código institucional."
            )

        else:

            ahora = hora_peru()

            reg = leer_registro(
                clase["id"],
                codigo
            )

            eventos = []

            # ------------------------------------------------
            # PRIMER INGRESO
            # ------------------------------------------------

            if reg is None:

                reg = logica.nuevo_registro(
                    codigo,
                    ahora,
                    logica.a_fecha(
                        clase["hora_inicio"]
                    ),
                    clase["tolerancia_min"]
                )

            # ------------------------------------------------
            # VOLVIÓ A ENTRAR
            # ------------------------------------------------

            else:

                logica.revisar_tiempos(
                    reg,
                    ahora,
                    SEGUNDOS_SIN_SENAL,
                    eventos
                )

                if (
                    reg["desconectado_desde"]
                    and not reg["retiro"]
                ):

                    logica.registrar_reingreso(
                        reg,
                        ahora,
                        eventos
                    )

                if not reg["retiro"]:

                    reg["ultima_senal"] = (
                        ahora.isoformat(
                            timespec="seconds"
                        )
                    )

            guardar_registro(
                clase["id"],
                reg,
                eventos
            )

            session["codigo"] = codigo

            return redirect(
                url_for("aula")
            )

    return render_template(
        "ingreso.html",
        clase=clase,
        mensaje=mensaje
    )


# ============================================================
# AULA DEL ALUMNO
# ============================================================

@app.route("/aula")
def aula():

    clase = clase_activa()
    codigo = session.get("codigo")

    if not clase or not codigo:
        return redirect(
            url_for("ingreso")
        )

    reg = leer_registro(
        clase["id"],
        codigo
    )

    if reg is None:
        return redirect(
            url_for("ingreso")
        )

    estudiante = bd().execute(
        """
        SELECT *
        FROM estudiantes
        WHERE codigo = ?
        """,
        (codigo,)
    ).fetchone()

    return render_template(
        "aula.html",
        clase=clase,
        reg=reg,
        estudiante=estudiante,
        estado=logica.estado_final(reg)
    )


# ============================================================
# SEÑAL DEL NAVEGADOR
# ============================================================

@app.route("/api/senal", methods=["POST"])
def senal():

    clase = clase_activa()
    codigo = session.get("codigo")

    if not clase or not codigo:

        return jsonify(
            clase_activa=False
        )

    reg = leer_registro(
        clase["id"],
        codigo
    )

    if reg is None:

        return jsonify(
            clase_activa=False
        )

    datos = request.get_json(
        silent=True
    ) or {}

    ahora = hora_peru()
    eventos = []

    visible = datos.get(
        "visible",
        True
    )

    # --------------------------------------------------------
    # Comprobar si había una desconexión
    # --------------------------------------------------------

    logica.revisar_tiempos(
        reg,
        ahora,
        SEGUNDOS_SIN_SENAL,
        eventos
    )

    # --------------------------------------------------------
    # PÁGINA VISIBLE
    # --------------------------------------------------------

    if visible:

        if (
            reg["desconectado_desde"]
            and not reg["retiro"]
        ):

            logica.registrar_reingreso(
                reg,
                ahora,
                eventos
            )

        if not reg["retiro"]:

            reg["ultima_senal"] = (
                ahora.isoformat(
                    timespec="seconds"
                )
            )

    # --------------------------------------------------------
    # PÁGINA NO VISIBLE
    # --------------------------------------------------------

    else:

        logica.registrar_salida(
            reg,
            ahora,
            eventos
        )

        if not reg["retiro"]:

            reg["ultima_senal"] = (
                ahora.isoformat(
                    timespec="seconds"
                )
            )

    guardar_registro(
        clase["id"],
        reg,
        eventos
    )

    return jsonify(
        clase_activa=True,
        estado=logica.estado_final(reg),
        desconexiones=reg["desconexiones"]
    )


# ============================================================
# PANEL DOCENTE
# ============================================================

@app.route("/docente", methods=["GET", "POST"])
def docente():

    if (
        request.method == "POST"
        and "clave" in request.form
    ):

        if request.form["clave"] == CLAVE_DOCENTE:

            session["docente"] = True

        else:

            return render_template(
                "docente_login.html",
                mensaje="Clave incorrecta."
            )

    if not session.get("docente"):

        return render_template(
            "docente_login.html",
            mensaje=None
        )

    clase = clase_activa()
    filas = []

    if clase:

        revisar_clase(
            clase["id"]
        )

        filas = filas_reporte(
            clase["id"]
        )

    ultima = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE activa = 0
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()

    return render_template(
        "docente.html",
        clase=clase,
        filas=filas,
        ultima=ultima
    )


# ============================================================
# INICIAR CLASE
# ============================================================

@app.post("/docente/iniciar")
def iniciar_clase():

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    if not clase_activa():

        curso = (
            request.form.get("curso")
            or "Clase virtual"
        )

        tolerancia = int(
            request.form.get(
                "tolerancia"
            ) or 10
        )

        bd().execute(
            """
            INSERT INTO clases
            (
                curso,
                hora_inicio,
                tolerancia_min
            )
            VALUES (?, ?, ?)
            """,
            (
                curso,
                hora_peru().isoformat(
                    timespec="seconds"
                ),
                tolerancia
            )
        )

        bd().commit()

    return redirect(
        url_for("docente")
    )


# ============================================================
# FINALIZAR CLASE
# ============================================================

@app.post("/docente/finalizar")
def finalizar_clase():

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = clase_activa()

    if clase:

        revisar_clase(
            clase["id"]
        )

        bd().execute(
            """
            UPDATE clases
            SET activa = 0,
                hora_fin = ?
            WHERE id = ?
            """,
            (
                hora_peru().isoformat(
                    timespec="seconds"
                ),
                clase["id"]
            )
        )

        bd().commit()

        return redirect(
            url_for(
                "reporte",
                clase_id=clase["id"]
            )
        )

    return redirect(
        url_for("docente")
    )


# ============================================================
# GENERAR FILAS DEL REPORTE
# ============================================================

def filas_reporte(clase_id):

    """
    IMPORTANTE:

    Se recorren TODOS los estudiantes.

    Si el alumno nunca ingresó:
        estado_final = FALTA

    Si ingresó:
        se utiliza su registro real.
    """

    filas = []

    estudiantes = bd().execute(
        """
        SELECT *
        FROM estudiantes
        ORDER BY nombre
        """
    ).fetchall()

    for est in estudiantes:

        codigo = est["codigo"]

        reg = leer_registro(
            clase_id,
            codigo
        )

        # ====================================================
        # NUNCA INGRESÓ
        # ====================================================

        if reg is None:

            filas.append(
                {
                    "codigo": codigo,
                    "nombre": est["nombre"],
                    "hora_ingreso": "",
                    "estado_ingreso": "",
                    "desconexiones": 0,
                    "salidas": "",
                    "reingresos": "",
                    "conectado": False,
                    "estado_final": logica.FALTA
                }
            )

            continue

        # ====================================================
        # SÍ INGRESÓ
        # ====================================================

        eventos = bd().execute(
            """
            SELECT tipo, hora
            FROM eventos
            WHERE clase_id = ?
            AND codigo = ?
            ORDER BY hora
            """,
            (
                clase_id,
                codigo
            )
        ).fetchall()

        salidas = ", ".join(
            e["hora"][11:]
            for e in eventos
            if e["tipo"] == "salida"
        )

        reingresos = ", ".join(
            e["hora"][11:]
            for e in eventos
            if e["tipo"] == "reingreso"
        )

        conectado = (
            not reg["desconectado_desde"]
            and not reg["retiro"]
        )

        filas.append(
            {
                "codigo": codigo,
                "nombre": est["nombre"],
                "hora_ingreso": (
                    reg["hora_ingreso"][11:]
                    if reg["hora_ingreso"]
                    else ""
                ),
                "estado_ingreso": (
                    reg["estado_ingreso"]
                    or ""
                ),
                "desconexiones": (
                    reg["desconexiones"]
                    or 0
                ),
                "salidas": salidas,
                "reingresos": reingresos,
                "conectado": conectado,
                "estado_final": logica.estado_final(reg)
            }
        )

    return filas


# ============================================================
# REPORTE HTML
# ============================================================

@app.route("/reporte/<int:clase_id>")
def reporte(clase_id):

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE id = ?
        """,
        (clase_id,)
    ).fetchone()

    if not clase:

        return redirect(
            url_for("docente")
        )

    revisar_clase(
        clase_id
    )

    filas = filas_reporte(
        clase_id
    )

    estados = (
        logica.PRESENTE,
        logica.TARDANZA,
        logica.FALTA,
        logica.RETIRO_FALTA
    )

    resumen = {
        estado: sum(
            1
            for f in filas
            if f["estado_final"] == estado
        )
        for estado in estados
    }

    return render_template(
        "reporte.html",
        clase=clase,
        filas=filas,
        resumen=resumen
    )


# ============================================================
# EXPORTAR A EXCEL
# ============================================================

@app.route("/reporte/<int:clase_id>/excel")
def reporte_excel(clase_id):

    if not session.get("docente"):

        return redirect(
            url_for("docente")
        )

    clase = bd().execute(
        """
        SELECT *
        FROM clases
        WHERE id = ?
        """,
        (clase_id,)
    ).fetchone()

    if not clase:

        return redirect(
            url_for("docente")
        )

    revisar_clase(
        clase_id
    )

    filas = filas_reporte(
        clase_id
    )

    colores = {
        logica.PRESENTE: "C8E6D4",
        logica.TARDANZA: "F6E2B3",
        logica.FALTA: "F2C4C0",
        logica.RETIRO_FALTA: "DCCDEE"
    }

    libro = Workbook()

    hoja = libro.active
    hoja.title = "Asistencia"

    # --------------------------------------------------------
    # ENCABEZADO
    # --------------------------------------------------------

    hoja.append(
        [
            f"Reporte de asistencia - "
            f"{clase['curso']}"
        ]
    )

    hoja.append(
        [
            f"Fecha: {clase['hora_inicio'][:10]}  "
            f"Inicio: {clase['hora_inicio'][11:]}  "
            f"Fin: {(clase['hora_fin'] or '')[11:]}"
        ]
    )

    hoja.append([])

    hoja.append(
        [
            "Código",
            "Estudiante",
            "Hora de ingreso",
            "Estado de ingreso",
            "Desconexiones",
            "Horas de salida",
            "Horas de reingreso",
            "Estado final"
        ]
    )

    # --------------------------------------------------------
    # FORMATO DEL ENCABEZADO
    # --------------------------------------------------------

    for celda in hoja[4]:

        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celda.fill = PatternFill(
            "solid",
            fgColor="1F3A5F"
        )

    # --------------------------------------------------------
    # DATOS
    # --------------------------------------------------------

    for f in filas:

        estado = f["estado_final"]

        hoja.append(
            [
                f["codigo"],
                f["nombre"],
                f["hora_ingreso"],
                f["estado_ingreso"],
                f["desconexiones"],
                f["salidas"],
                f["reingresos"],
                estado
            ]
        )

        # Color según estado
        hoja.cell(
            hoja.max_row,
            8
        ).fill = PatternFill(
            "solid",
            fgColor=colores.get(
                estado,
                "FFFFFF"
            )
        )

    # --------------------------------------------------------
    # ANCHOS
    # --------------------------------------------------------

    hoja["A1"].font = Font(
        bold=True,
        size=14
    )

    anchos = (
        12,
        32,
        15,
        17,
        14,
        24,
        24,
        16
    )

    for col, ancho in zip(
        "ABCDEFGH",
        anchos
    ):

        hoja.column_dimensions[
            col
        ].width = ancho

    # --------------------------------------------------------
    # CREAR ARCHIVO
    # --------------------------------------------------------

    archivo = io.BytesIO()

    libro.save(
        archivo
    )

    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name=(
            f"asistencia_"
            f"{clase['hora_inicio'][:10]}_"
            f"clase{clase_id}.xlsx"
        )
    )


# ============================================================
# SALIR
# ============================================================

@app.route("/salir")
def salir():

    session.clear()

    return redirect(
        url_for("ingreso")
    )


# ============================================================
# CREAR BASE DE DATOS
# ============================================================

crear_tablas()


# ============================================================
# EJECUTAR
# ============================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )
